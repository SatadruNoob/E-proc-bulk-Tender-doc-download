#!/usr/bin/env python3
"""
Semi-automated tender document downloader for GePNIC (NIC e-procurement) portals.

What it automates : reading your Excel, opening the right portal, searching the Tender ID on the
                    home-page keyword search (no captcha there), opening the tender, clicking
                    "Download as zip", naming/sorting the file, logging, resuming.
What YOU do       : type the captcha on the download page when prompted (a few seconds each).

Setup (Windows, once, inside your venv):
    python -m pip install playwright openpyxl
    python -m playwright install chromium

Usage:
    python download_tenders.py download_tender_doc.xlsx --check-portals   # does every portal open?
    python download_tenders.py download_tender_doc.xlsx --dry-run         # show the plan
    python download_tenders.py download_tender_doc.xlsx --limit 3         # try 3 first
    python download_tenders.py download_tender_doc.xlsx --only-new        # only "New Today" = YES
    python download_tenders.py download_tender_doc.xlsx --state Kerala    # one portal only

Files land in  downloads/<State>/<TenderID>.zip   and results in  download_log.csv
Re-running skips tenders already downloaded.
"""
import argparse
import csv
import json
import re
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# Same portal list as your scraper (keys must match the "State/UT" column in the sheet).
PORTALS = {
    "Andaman": "https://eprocure.andamannicobar.gov.in/nicgep/app",
    "Arunachal": "https://arunachaltenders.gov.in/nicgep/app",
    "Assam": "https://assamtenders.gov.in/nicgep/app",
    "Chandigarh": "https://etenders.chd.nic.in/nicgep/app",
    "UT Nagar": "https://dnhtenders.gov.in/nicgep/app",
    "UT Daman": "https://ddtenders.gov.in/nicgep/app",
    "NCT Delhi": "https://govtprocurement.delhi.gov.in/nicgep/app",
    "Goa": "https://eprocure.goa.gov.in/nicgep/app",
    "Haryana": "https://etenders.hry.nic.in/nicgep/app",
    "Himachal": "https://hptenders.gov.in/nicgep/app",
    "Jammu": "https://jktenders.gov.in/nicgep/app",
    "Jharkhand": "https://jharkhandtenders.gov.in/nicgep/app",
    "Kerala": "https://etenders.kerala.gov.in/nicgep/app",
    "Lakshadweep": "https://tendersutl.gov.in/nicgep/app",
    "Maharashtra": "https://mahatenders.gov.in/nicgep/app",
    "MP": "https://mptenders.gov.in/nicgep/app",
    "Manipur": "https://manipurtenders.gov.in/nicgep/app",
    "Meghalaya": "https://meghalayatenders.gov.in/nicgep/app",
    "Mizoram": "https://mizoramtenders.gov.in/nicgep/app",
    "Nagaland": "https://nagalandtenders.gov.in/nicgep/app",
    "Odisha": "https://tendersodisha.gov.in/nicgep/app",
    "Puducherry": "https://pudutenders.gov.in/nicgep/app",
    "Punjab": "https://eproc.punjab.gov.in/nicgep/app",
    "Rajasthan": "https://eproc.rajasthan.gov.in/nicgep/app",
    "Sikkim": "https://sikkimtender.gov.in/nicgep/app",
    "Tamil Nadu": "https://tntenders.gov.in/nicgep/app",
    "Tripura": "https://tripuratenders.gov.in/nicgep/app",
    "Ladakh": "https://tenders.ladakh.gov.in/nicgep/app",
    "Uttarakhand": "https://uktenders.gov.in/nicgep/app",
    "Uttar Pradesh": "https://etender.up.nic.in/nicgep/app",
    "West Bengal": "https://wbtenders.gov.in/nicgep/app",
    "Central Govt eProcurement": "https://eprocure.gov.in/eprocure/app",
    "CPSE eProcurement": "https://etenders.gov.in/eprocure/app",
    "NTPC Limited": "https://eprocurentpc.nic.in/nicgep/app",
    "MIDHANI": "https://eprocuremidhani.nic.in/nicgep/app",
    "Mazagon Dock Shipbuilders": "https://eprocuremdl.nic.in/nicgep/app",
    "IOCL": "https://iocletenders.nic.in/nicgep/app",
    "Hindustan Shipyard Limited": "https://eprocurehsl.nic.in/nicgep/app",
    "Goa Shipyard Limited": "https://eprocuregsl.nic.in/nicgep/app",
    "Garden Reach Shipbuilders (GRSE)": "https://eprocuregrse.co.in/nicgep/app",
    "CPCL": "https://cpcletenders.nic.in/nicgep/app",
    "Coal India Limited": "https://coalindiatenders.nic.in/nicgep/app",
    "BHEL": "https://eprocurebhel.co.in/nicgep/app",
    "BEL Defence": "https://eprocurebel.co.in/nicgep/app",
    "PMGSY": "https://pmgsytenders.gov.in/nicgep/app",
    "Defence e-Procurement": "https://defproc.gov.in/nicgep/app",
}

# Optional: portals.json next to this script overrides/adds portals, e.g. {"Assam": "https://..."}
_override = Path("portals.json")
if _override.exists():
    PORTALS.update(json.loads(_override.read_text(encoding="utf-8")))

SEARCH_FALLBACK = "batter"      # same keyword as your scraper; used if a Tender ID search finds nothing
OUT_DIR = Path("downloads")
LOG_FILE = Path("download_log.csv")


def parse_closing(value):
    """'12-Oct-2026 03:00 PM' -> datetime (None if it can't be read)."""
    if isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%d-%b-%Y %I:%M %p")
    except ValueError:
        return None


def read_tenders(xlsx):
    ws = load_workbook(xlsx, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    ix = {h: i for i, h in enumerate(header)}
    for need in ("State/UT", "Tender ID"):
        if need not in ix:
            raise SystemExit(f"Column '{need}' not found in sheet. Found: {header}")
    tenders, seen = [], set()
    for r in rows:
        tid = r[ix["Tender ID"]]
        if not tid or str(tid).strip() in seen:
            continue
        seen.add(str(tid).strip())
        tenders.append({
            "state": str(r[ix["State/UT"]]).strip(),
            "tender_id": str(tid).strip(),
            "new": ("New Today" not in ix) or str(r[ix["New Today"]]).strip().upper() == "YES",
            "closing": parse_closing(r[ix["Closing Date"]]) if "Closing Date" in ix else None,
        })
    return tenders


def folder_for(state):
    return OUT_DIR / re.sub(r"[^\w\- ]", "_", state)


def already_downloaded(t):
    folder = folder_for(t["state"])
    return any(folder.glob(f"{t['tender_id']}.*")) or (folder / t["tender_id"]).is_dir()


def log(t, status, note=""):
    new = not LOG_FILE.exists()
    with LOG_FILE.open("a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["time", "state", "tender_id", "status", "note"])
        w.writerow([datetime.now().isoformat(timespec="seconds"), t["state"], t["tender_id"], status, note])


PROBE_DIR = Path("probe")


def dump_page(page, tag):
    """Save URL, all links/buttons, HTML and a screenshot of a page so its structure can be inspected."""
    PROBE_DIR.mkdir(exist_ok=True)
    base = PROBE_DIR / tag
    try:
        items = page.eval_on_selector_all(
            "a, button, input[type=submit], input[type=button]",
            """els => els.map(e => ({tag: e.tagName, id: e.id, name: e.getAttribute('name'),
                 text: (e.innerText || e.value || '').trim().slice(0, 100),
                 href: (e.getAttribute('href') || '').slice(0, 200)}))""")
        lines = [page.url, f"{len(items)} links/buttons", ""] + [json.dumps(i, ensure_ascii=False) for i in items]
        base.with_suffix(".txt").write_text("\n".join(lines), encoding="utf-8")
        base.with_suffix(".html").write_text(page.content(), encoding="utf-8")
        page.screenshot(path=str(base.with_suffix(".png")), full_page=True)
        print(f"  [probe] saved {base}.txt/.html/.png")
    except Exception as e:
        print(f"  [probe] could not dump {tag}: {str(e)[:100]}")


def open_portal(page, base):
    """Open the portal home page and dismiss the pop-up alert if present."""
    page.goto(base, wait_until="load")
    if page.locator("button.alertbutclose").count() > 0:
        page.click("button.alertbutclose")
        page.wait_for_timeout(800)


def search_for_row(page, base, tender_id):
    """Keyword-search (no captcha) and return the result row containing tender_id, or None."""
    for term in (tender_id, SEARCH_FALLBACK):
        open_portal(page, base)
        page.fill("#SearchDescription", term)
        page.click("#Go")
        page.wait_for_load_state("networkidle")
        row = page.locator("tr.even, tr.odd").filter(has_text=tender_id)
        if row.count() > 0:
            return row.first
    return None


def process(page, context, downloads, t, timeout, settle, probe=False):
    from playwright.sync_api import TimeoutError as PWTimeout

    base = PORTALS[t["state"]]
    downloads.clear()
    for p in context.pages:                      # close leftover tabs from last tender
        if p is not page:
            p.close()
    try:
        page.goto("about:blank")                 # cancel any navigation left over from a failed portal
    except Exception:
        pass

    # Step 1: find the tender via the home-page keyword search (no captcha) and open it
    row = search_for_row(page, base, t["tender_id"])
    if row is None:
        return "not_found", "not in the portal's active listing (closed/withdrawn?)"
    row.locator("td").nth(4).locator("a").first.click()
    page.wait_for_load_state("load")
    if probe:
        dump_page(page, f"{t['tender_id']}_1_detail")

    # Step 2: open the zip download (captcha is on the next page)
    zip_link = page.get_by_text(re.compile(r"download as zip", re.I)).first
    try:
        zip_link.wait_for(timeout=30_000)
        zip_link.click()
        print("  -> Type the captcha in the browser and click Submit")
        if probe:
            page.wait_for_timeout(2000)
            dump_page(context.pages[-1], f"{t['tender_id']}_2_captcha_page")
    except PWTimeout:
        print("  -> Zip link not found. Click the download link yourself in the browser")

    # Step 3: wait for the file (works even if you clicked things manually)
    end = time.time() + timeout
    while time.time() < end and not downloads:
        page.wait_for_timeout(500)
    if not downloads:
        return "timeout", "no download captured"
    if probe:
        page.wait_for_timeout(3000)
        dump_page(context.pages[-1], f"{t['tender_id']}_3_after_download")

    # One captcha can release several files (or you may click extra attachments yourself).
    # Keep collecting until no new file has arrived for `settle` seconds.
    last_count, last_change = len(downloads), time.time()
    while time.time() - last_change < settle:
        page.wait_for_timeout(500)
        if len(downloads) != last_count:
            last_count, last_change = len(downloads), time.time()

    files = list(downloads)
    folder = folder_for(t["state"])
    folder.mkdir(parents=True, exist_ok=True)

    if len(files) == 1:                                   # usual case: <TenderID>.zip / .pdf
        ext = Path(files[0].suggested_filename).suffix or ".zip"
        dest = folder / f"{t['tender_id']}{ext}"
        files[0].save_as(dest)
        return "ok", f"{dest} (portal file name: {files[0].suggested_filename})"

    sub = folder / t["tender_id"]                         # several files: <TenderID>/<original names>
    sub.mkdir(exist_ok=True)
    for n, d in enumerate(files, 1):
        name = re.sub(r'[\\/:*?"<>|]', "_", d.suggested_filename or f"file_{n}")
        dest = sub / name
        if dest.exists():
            dest = sub / f"{dest.stem}_{n}{dest.suffix}"
        d.save_as(dest)
    return "ok", f"{len(files)} files in {sub}: " + "; ".join(d.suggested_filename for d in files)


def check_portals(ignore_cert):
    """Load each distinct portal once and report which ones open."""
    from playwright.sync_api import sync_playwright

    seen = {}
    for name, url in PORTALS.items():
        seen.setdefault(url, []).append(name)
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=False)
        page = browser.new_context(ignore_https_errors=ignore_cert).new_page()
        for url, names in seen.items():
            try:
                resp = page.goto(url, wait_until="domcontentloaded", timeout=30_000)
                has_search = page.locator("#SearchDescription").count() > 0
                print(f"OK   {resp.status if resp else '?'} search-box={'yes' if has_search else 'NO'}  "
                      f"{url}   ({', '.join(names)})")
            except Exception as e:
                print(f"FAIL {url}   ({', '.join(names)})  -> {str(e).splitlines()[0][:100]}")
            try:
                page.goto("about:blank")
            except Exception:
                pass
        browser.close()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--only-new", action="store_true", help="only rows with New Today = YES")
    ap.add_argument("--state", help="only this State/UT value, e.g. Kerala")
    ap.add_argument("--limit", type=int, help="stop after N tenders (good for a first test)")
    ap.add_argument("--timeout", type=int, default=180, help="seconds to wait for you per tender (default 180)")
    ap.add_argument("--include-closed", action="store_true",
                    help="also try tenders whose Closing Date has passed (they are normally gone from the portals)")
    ap.add_argument("--settle", type=int, default=6,
                    help="after the first file, wait this many seconds for more files from the same captcha (default 6)")
    ap.add_argument("--tender-id", help="only this Tender ID")
    ap.add_argument("--probe", action="store_true",
                    help="save the page structure (links, HTML, screenshot) at each step into probe/ for inspection")
    ap.add_argument("--dry-run", action="store_true", help="show the plan and exit")
    ap.add_argument("--check-portals", action="store_true", help="test that every portal URL opens, then exit")
    ap.add_argument("--ignore-cert-errors", action="store_true", help="for portals with SSL certificate problems")
    args = ap.parse_args()

    if args.check_portals:
        check_portals(args.ignore_cert_errors)
        return

    tenders = read_tenders(args.xlsx)
    if args.only_new:
        tenders = [t for t in tenders if t["new"]]
    if args.state:
        tenders = [t for t in tenders if t["state"].lower() == args.state.lower()]
    if args.tender_id:
        tenders = [t for t in tenders if t["tender_id"] == args.tender_id]

    now = datetime.now()
    closed = [t for t in tenders if not args.include_closed and t.get("closing") and t["closing"] < now]
    if closed:
        print(f"Skipping {len(closed)} tenders whose closing date has passed (use --include-closed to try them)")
        tenders = [t for t in tenders if t not in closed]

    todo = [t for t in tenders if t["state"] in PORTALS and (args.probe or not already_downloaded(t))]
    unknown = sorted({t["state"] for t in tenders if t["state"] not in PORTALS})
    todo.sort(key=lambda t: (t["state"], t.get("closing") or datetime.max))   # same portal back to back
    if args.limit:
        todo = todo[: args.limit]

    print(f"{len(tenders)} tenders in scope | {len(todo)} to download | "
          f"{sum(already_downloaded(t) for t in tenders)} already done")
    if unknown:
        print(f"No portal URL configured for: {unknown}  (add them to portals.json)")
    per_portal = {}
    for t in todo:
        per_portal[t["state"]] = per_portal.get(t["state"], 0) + 1
    print("Per portal:", ", ".join(f"{k} {v}" for k, v in per_portal.items()))
    if args.dry_run:
        for t in todo:
            print(f"  {t['state']:28} {t['tender_id']:26} {PORTALS[t['state']]}")
        return

    from playwright.sync_api import sync_playwright

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=False)
        context = browser.new_context(accept_downloads=True, ignore_https_errors=args.ignore_cert_errors)
        downloads = []

        def on_download(d):
            downloads.append(d)

        def on_new_page(p):
            p.on("download", on_download)

        context.on("page", on_new_page)
        page = context.new_page()   # fires on_new_page, so downloads on every tab are captured
        page.set_default_timeout(90_000)

        try:
            for i, t in enumerate(todo, 1):
                print(f"[{i}/{len(todo)}] {t['state']} | {t['tender_id']}")
                try:
                    status, note = process(page, context, downloads, t, args.timeout, args.settle, args.probe)
                except Exception as e:  # keep going on any portal hiccup
                    status, note = "error", str(e)[:200]
                print(f"  = {status} {note}")
                log(t, status, note)
        except KeyboardInterrupt:
            print("\nStopped. Re-run to resume; finished tenders are skipped.")
        finally:
            browser.close()


if __name__ == "__main__":
    main()
